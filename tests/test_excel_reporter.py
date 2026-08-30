import pytest
import os
import openpyxl
from src.data_generator import generate_recruitment_dataset
from src.tableau_prep_engine import TableauPrepEngine
from src.excel_reporter import ExecutiveExcelReporter

@pytest.fixture(scope="module")
def prepared_fact_file(tmp_path_factory):
    raw_fn = tmp_path_factory.mktemp("raw") / "raw.parquet"
    df = generate_recruitment_dataset(num_records=1000, random_seed=42)
    df.to_parquet(str(raw_fn), index=False)

    engine = TableauPrepEngine()
    df_fact = engine.execute_star_schema(raw_parquet_path=str(raw_fn))
    
    fact_fn = tmp_path_factory.mktemp("curated") / "fact.parquet"
    df_fact.to_parquet(str(fact_fn), index=False)
    return str(fact_fn)

def test_executive_excel_generation(prepared_fact_file, tmp_path):
    out_xlsx = tmp_path / "test_scorecard.xlsx"
    reporter = ExecutiveExcelReporter(parquet_path=prepared_fact_file)
    path = reporter.generate_executive_workbook(output_path=str(out_xlsx))

    assert os.path.exists(path)
    assert os.path.getsize(path) > 5000

    # Inspect Workbook
    wb = openpyxl.load_workbook(path)
    assert "Executive Summary" in wb.sheetnames
    ws = wb["Executive Summary"]
    assert "APPLY ON JOB" in ws["A1"].value
